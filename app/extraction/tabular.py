"""Import of structured invoice-line and meter-reading tables (CSV / XLSX).

Canonical format: one row per invoice line; invoice header columns repeat per row.
Dutch and English column names are accepted (see COLUMN_ALIASES). See docs/import-formats.md.
"""

from __future__ import annotations

import csv
import io
import re
from collections import OrderedDict
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from app.domain.ean import is_valid_ean18, normalize_ean
from app.domain.enums import Commodity, ExtractionMethod, InvoiceType, LineCategory, ReadingType
from app.domain.units import normalize_unit
from app.extraction.categories import classify_line
from app.extraction.invoice_pdf import validate_invoice
from app.extraction.parsing import parse_date, parse_decimal
from app.extraction.schema import ExtractedInvoice, ExtractedLine, ExtractedReading, ExtractionResult, Field

COLUMN_ALIASES: dict[str, list[str]] = {
    "invoice_number": ["factuurnummer", "factuurnr", "invoice_number", "invoice number", "notanummer"],
    "invoice_date": ["factuurdatum", "invoice_date", "invoice date", "datum factuur"],
    "supplier": ["leverancier", "supplier", "energieleverancier"],
    "ean": ["ean", "ean_code", "ean-code", "eancode", "aansluiting", "ean code"],
    "meter_number": ["meternummer", "meter_number", "meter", "meternr"],
    "billing_period_start": ["periode_start", "period_start", "periode van", "van", "begindatum", "startdatum"],
    "billing_period_end": ["periode_eind", "period_end", "periode tot", "tot", "t/m", "einddatum", "tm"],
    "invoice_type": ["factuurtype", "invoice_type", "type factuur", "soort"],
    "corrects_invoice_number": ["correctie_op", "corrects_invoice_number", "originele factuur", "betreft factuur"],
    "commodity": ["product", "commodity", "energiesoort"],
    "description": ["omschrijving", "description", "regel", "kostenpost"],
    "category": ["categorie", "category"],
    "quantity": ["hoeveelheid", "aantal", "quantity", "verbruik", "volume"],
    "unit": ["eenheid", "unit"],
    "unit_price": ["tarief", "prijs", "unit_price", "eenheidsprijs", "prijs per eenheid"],
    "amount": ["bedrag", "amount", "bedrag_excl_btw", "bedrag excl btw", "regelbedrag"],
    "vat_rate": ["btw_percentage", "vat_rate", "btw%", "btw %", "btw tarief"],
    "line_period_start": ["regel_van", "line_period_start"],
    "line_period_end": ["regel_tot", "line_period_end"],
    "subtotal_excl_vat": ["subtotaal", "subtotal", "totaal_excl_btw", "totaal excl btw"],
    "vat_amount": ["btw_bedrag", "vat_amount", "btw bedrag", "btw"],
    "total_incl_vat": ["totaal", "total", "totaal_incl_btw", "totaal incl btw", "factuurbedrag"],
}
READING_ALIASES: dict[str, list[str]] = {
    "ean": COLUMN_ALIASES["ean"],
    "meter_number": COLUMN_ALIASES["meter_number"],
    "register": ["register", "telwerk", "tariefperiode"],
    "reading_date": ["datum", "reading_date", "opnamedatum", "date"],
    "reading": ["stand", "meterstand", "reading", "value"],
    "reading_type": ["type", "reading_type", "soort opname", "opnametype"],
    "multiplier": ["vermenigvuldigingsfactor", "multiplier", "factor"],
    "unit": ["eenheid", "unit"],
}

INVOICE_REQUIRED = {"invoice_number", "amount"}
READING_REQUIRED = {"reading_date", "reading"}


def _norm_header(h: Any) -> str:
    return re.sub(r"\s+", " ", str(h or "").strip().lower().replace("_", " ")).replace(" ", "_")


def _map_headers(headers: list[Any], aliases: dict[str, list[str]]) -> dict[str, int]:
    normalized = [_norm_header(h) for h in headers]
    mapping: dict[str, int] = {}
    for canonical, names in aliases.items():
        wanted = {_norm_header(n) for n in [canonical, *names]}
        for idx, h in enumerate(normalized):
            if h in wanted and idx not in mapping.values():
                mapping[canonical] = idx
                break
    return mapping


def read_rows(content: bytes, kind: str) -> list[list[Any]]:
    if kind == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                rows = [list(r) for r in ws.iter_rows(values_only=True)]
                if any(any(c is not None for c in r) for r in rows):
                    return rows
            return []
        finally:
            wb.close()
    text = _decode(content)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";" if sample.count(";") > sample.count(",") else ","
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def _decode(content: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ValueError("Onbekende tekencodering")


def detect_table_kind(rows: list[list[Any]]) -> tuple[str | None, int, dict[str, int]]:
    """Find header row and decide whether this is an invoice table or meter data."""
    for idx, row in enumerate(rows[:15]):
        inv = _map_headers(row, COLUMN_ALIASES)
        if INVOICE_REQUIRED <= inv.keys() and len(inv) >= 4:
            return "invoice", idx, inv
        rd = _map_headers(row, READING_ALIASES)
        if READING_REQUIRED <= rd.keys():
            return "meter", idx, rd
    return None, 0, {}


def _cell(row: list[Any], mapping: dict[str, int], key: str) -> Any:
    idx = mapping.get(key)
    if idx is None or idx >= len(row):
        return None
    v = row[idx]
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def _dec(value: Any) -> tuple[Decimal | None, bool]:
    if value is None:
        return None, False
    if isinstance(value, bool):
        return None, False
    if isinstance(value, int):
        return Decimal(value), False
    if isinstance(value, float):
        # Spreadsheet numeric cells arrive as binary floats: use the shortest repr, which
        # equals what the spreadsheet displays for values typed by humans.
        return Decimal(repr(value)), False
    if isinstance(value, Decimal):
        return value, False
    parsed = parse_decimal(str(value))
    return (parsed.value, parsed.ambiguous) if parsed else (None, False)


def _date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return parse_date(str(value)) if value is not None else None


def parse_invoice_table(rows: list[list[Any]], header_idx: int, mapping: dict[str, int],
                        method: ExtractionMethod) -> ExtractionResult:
    invoices: OrderedDict[str, ExtractedInvoice] = OrderedDict()
    notes: list[str] = []
    for offset, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if not any(c not in (None, "") for c in row):
            continue
        number = _cell(row, mapping, "invoice_number")
        if number is None:
            notes.append(f"Rij {offset}: geen factuurnummer — overgeslagen.")
            continue
        number = str(number).strip()
        inv = invoices.get(number)
        if inv is None:
            inv = ExtractedInvoice(source_row=offset)
            invoices[number] = inv
            _table_header(inv, row, mapping, offset, method, number)

        desc = _cell(row, mapping, "description") or ""
        qty, qa = _dec(_cell(row, mapping, "quantity"))
        price, pa = _dec(_cell(row, mapping, "unit_price"))
        amount, aa = _dec(_cell(row, mapping, "amount"))
        vat_rate, _ = _dec(_cell(row, mapping, "vat_rate"))
        unit = normalize_unit(_cell(row, mapping, "unit"))
        explicit_cat = _cell(row, mapping, "category")
        if explicit_cat and str(explicit_cat).upper() in LineCategory.__members__:
            category, cat_conf = LineCategory(str(explicit_cat).upper()), 1.0
        else:
            category, cat_conf = classify_line(str(desc), unit)
        if amount is None and qty is None:
            notes.append(f"Rij {offset}: geen bedrag of hoeveelheid — overgeslagen.")
            continue
        conf = 0.98 - (0.1 if (qa or pa or aa) else 0.0)
        inv.lines.append(ExtractedLine(
            description=str(desc), category=category, category_confidence=cat_conf,
            quantity=qty, unit=unit, unit_price=price, amount=amount, vat_rate=vat_rate,
            period_start=_date(_cell(row, mapping, "line_period_start")),
            period_end=_date(_cell(row, mapping, "line_period_end")),
            row=offset, source_text=" | ".join("" if c is None else str(c) for c in row),
            method=method, confidence=conf,
        ))

    for inv in invoices.values():
        _derive_commodity_from_lines(inv)
        validate_invoice(inv)
        # tabular sources are structured: missing EAN/lines are the main risks
    overall = min((i.confidence for i in invoices.values()), default=0.0)
    return ExtractionResult(invoices=list(invoices.values()), notes=notes, confidence=overall)


def _table_header(inv: ExtractedInvoice, row, mapping, offset, method, number) -> None:
    def put(name: str, value: Any, conf: float = 0.98) -> None:
        if value is not None:
            raw = _cell(row, mapping, name)
            inv.fields[name] = Field(value=value, raw_text=None if raw is None else str(raw), row=offset,
                                     confidence=conf, method=method)

    put("invoice_number", number)
    put("invoice_date", _date(_cell(row, mapping, "invoice_date")))
    supplier = _cell(row, mapping, "supplier")
    put("supplier", str(supplier) if supplier else None)
    ean = normalize_ean(str(_cell(row, mapping, "ean"))) if _cell(row, mapping, "ean") is not None else None
    if ean:
        put("ean", ean, 0.98 if is_valid_ean18(ean) else 0.5)
        if not is_valid_ean18(ean):
            inv.notes.append(f"EAN {ean} heeft een ongeldig controlegetal.")
    meter = _cell(row, mapping, "meter_number")
    put("meter_number", str(meter) if meter is not None else None)
    put("billing_period_start", _date(_cell(row, mapping, "billing_period_start")))
    put("billing_period_end", _date(_cell(row, mapping, "billing_period_end")))
    itype = str(_cell(row, mapping, "invoice_type") or "").lower()
    if itype:
        put("invoice_type", InvoiceType.CREDIT_NOTE if "credit" in itype else InvoiceType.INVOICE)
    corrects = _cell(row, mapping, "corrects_invoice_number")
    put("corrects_invoice_number", str(corrects) if corrects else None)
    commodity = str(_cell(row, mapping, "commodity") or "").lower()
    if commodity:
        put("commodity", Commodity.GAS if "gas" in commodity else
            Commodity.ELECTRICITY if ("elek" in commodity or "stroom" in commodity or "elec" in commodity)
            else Commodity.UNKNOWN)
    for name in ("subtotal_excl_vat", "vat_amount", "total_incl_vat"):
        put(name, _dec(_cell(row, mapping, name))[0])


def _derive_commodity_from_lines(inv: ExtractedInvoice) -> None:
    if inv.get("commodity") not in (None, Commodity.UNKNOWN):
        return
    from app.extraction.invoice_pdf import _derive_commodity

    _derive_commodity(inv)


_READING_TYPES = {
    "werkelijk": ReadingType.ACTUAL, "actual": ReadingType.ACTUAL, "gemeten": ReadingType.ACTUAL,
    "opgenomen": ReadingType.ACTUAL, "fysiek": ReadingType.ACTUAL, "geschat": ReadingType.ESTIMATED,
    "estimated": ReadingType.ESTIMATED, "berekend": ReadingType.ESTIMATED,
}


def parse_meter_table(rows: list[list[Any]], header_idx: int, mapping: dict[str, int]) -> ExtractionResult:
    readings: list[ExtractedReading] = []
    notes: list[str] = []
    for offset, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if not any(c not in (None, "") for c in row):
            continue
        d = _date(_cell(row, mapping, "reading_date"))
        value, _ = _dec(_cell(row, mapping, "reading"))
        if d is None or value is None:
            notes.append(f"Rij {offset}: ongeldige datum of stand — overgeslagen.")
            continue
        mult, _ = _dec(_cell(row, mapping, "multiplier"))
        rtype_raw = str(_cell(row, mapping, "reading_type") or "").lower()
        reg_raw = str(_cell(row, mapping, "register") or "").upper()
        register = {"NORMAAL": "NORMAL", "PIEK": "NORMAL", "DAL": "LOW", "LAAG": "LOW",
                    "ENKEL": "SINGLE", "TERUGLEVERING": "FEED_IN"}.get(reg_raw, reg_raw or "SINGLE")
        unit = normalize_unit(_cell(row, mapping, "unit")) or "kWh"
        if unit == "m3" and register == "SINGLE":
            register = "GAS"
        ean = normalize_ean(str(_cell(row, mapping, "ean"))) if _cell(row, mapping, "ean") is not None else None
        meter = _cell(row, mapping, "meter_number")
        readings.append(ExtractedReading(
            reading_date=d, reading=value, reading_type=_READING_TYPES.get(rtype_raw, ReadingType.UNKNOWN),
            register=register, ean=ean, meter_number=str(meter) if meter is not None else None,
            multiplier=mult or Decimal(1), unit=unit, row=offset,
            source_text=" | ".join("" if c is None else str(c) for c in row),
        ))
    return ExtractionResult(readings=readings, notes=notes, confidence=1.0 if readings else 0.0)


def parse_table(content: bytes, kind: str) -> tuple[str | None, ExtractionResult]:
    rows = read_rows(content, kind)
    table_kind, header_idx, mapping = detect_table_kind(rows)
    method = ExtractionMethod.XLSX if kind == "xlsx" else ExtractionMethod.CSV
    if table_kind == "invoice":
        return "invoice", parse_invoice_table(rows, header_idx, mapping, method)
    if table_kind == "meter":
        return "meter", parse_meter_table(rows, header_idx, mapping)
    return None, ExtractionResult(notes=["Kolomkoppen niet herkend. Zie het importformaat in de documentatie."])
